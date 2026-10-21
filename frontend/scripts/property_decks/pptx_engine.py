import json
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x00, 0x33, 0x66)
DARK_NAVY = RGBColor(0x10, 0x25, 0x3F)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)
WARM_WHITE = RGBColor(0xFF, 0xFC, 0xFB)
SLATE_TEXT = RGBColor(0x1E, 0x29, 0x3B)
MUTED_TEXT = RGBColor(0x64, 0x74, 0x8B)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
GOLD_TINT = RGBColor(0xFB, 0xF7, 0xE9)


def add_solid_bg(slide, prs, color):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def add_cover_slide(prs, proposal):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_solid_bg(slide, prs, NAVY)

    gold_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.12), Inches(4.8)
    )
    gold_bar.fill.solid()
    gold_bar.fill.fore_color.rgb = GOLD
    gold_bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.8), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PRIME PHILIPPINES  |  INDUSTRIAL & LOGISTICS ADVISORY"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = GOLD
    p0.font.name = "Century Gothic"

    p1 = tf.add_paragraph()
    p1.text = proposal.get("projectTitle") or "Warehouse & Industrial Property Proposal"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WARM_WHITE
    p1.font.name = "Century Gothic"
    p1.space_before = Pt(18)

    p2 = tf.add_paragraph()
    client_name = proposal.get("clientName") or "Valued Client"
    p2.text = f"Prepared Exclusively For: {client_name}"
    p2.font.size = Pt(18)
    p2.font.color.rgb = GOLD
    p2.font.name = "Century Gothic"
    p2.space_before = Pt(24)

    p3 = tf.add_paragraph()
    broker = proposal.get("brokerName") or "PRIME Industrial Brokerage"
    date_str = proposal.get("date") or ""
    p3.text = f"Prepared by: {broker}   |   {date_str}"
    p3.font.size = Pt(12)
    p3.font.color.rgb = WARM_WHITE
    p3.font.name = "Century Gothic"
    p3.space_before = Pt(40)


def add_header_banner(slide, prs, eyebrow, title):
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.05), prs.slide_width, Inches(0.05)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.14), Inches(12.0), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = eyebrow.upper()
    p0.font.size = Pt(9)
    p0.font.bold = True
    p0.font.color.rgb = GOLD
    p0.font.name = "Century Gothic"

    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = WARM_WHITE
    p1.font.name = "Century Gothic"


def add_summary_table_slide(prs, rows, proposal):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_solid_bg(slide, prs, WARM_WHITE)
    add_header_banner(
        slide,
        prs,
        "Executive Summary",
        f"Shortlisted Warehouse Options — {proposal.get('clientName') or 'Client Requirement'}",
    )

    headers = [
        "Option",
        "Property / Warehouse",
        "Location",
        "Leasable Area (GLA)",
        "Clear Height",
        "Lease Rate",
        "Remarks / Notes",
    ]
    num_rows = len(rows) + 1
    num_cols = len(headers)

    table_shape = slide.shapes.add_table(
        num_rows,
        num_cols,
        Inches(0.6),
        Inches(1.4),
        Inches(12.13),
        Inches(min(5.4, 0.55 * num_rows)),
    )
    table = table_shape.table

    col_widths = [
        Inches(0.8),
        Inches(2.7),
        Inches(2.0),
        Inches(1.6),
        Inches(1.3),
        Inches(1.6),
        Inches(2.13),
    ]
    for idx, width in enumerate(col_widths):
        table.columns[idx].width = width

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WARM_WHITE
        p.font.name = "Century Gothic"

    for row_idx, item in enumerate(rows, start=1):
        values = [
            f"Option {item.get('optionNumber', row_idx)}",
            str(item.get("propertyName") or ""),
            str(item.get("location") or ""),
            str(item.get("gla") or ""),
            str(item.get("ceilingHeight") or ""),
            str(item.get("leaseRate") or ""),
            str(item.get("notes") or "Standard handover"),
        ]
        bg_color = LIGHT_BG if row_idx % 2 == 0 else WARM_WHITE
        for col_idx, val in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.bold = col_idx in (0, 1)
            p.font.color.rgb = NAVY if col_idx in (0, 1) else SLATE_TEXT
            p.font.name = "Century Gothic"


def add_property_slide(prs, slide_item, option_idx, thumb_dir=None):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_solid_bg(slide, prs, WARM_WHITE)

    specs = slide_item.get("specs") or {}
    title = specs.get("propertyName") or slide_item.get("title") or f"Warehouse Option {option_idx}"
    location = specs.get("location") or "Industrial Corridor"

    add_header_banner(
        slide,
        prs,
        f"Option {option_idx}  ·  {location}",
        title,
    )

    spec_pairs = [
        ("Location", specs.get("location") or "Upon Request"),
        ("Gross Leasable Area", specs.get("gla") or "Upon Request"),
        ("Clear Ceiling Height", specs.get("ceilingHeight") or "Standard Industrial"),
        ("Asking Lease Rate", specs.get("leaseRate") or "Upon Request"),
        ("Power Supply", specs.get("powerSupply") or "3-Phase"),
        ("Loading Bays", specs.get("loadingBays") or "Available"),
    ]
    if specs.get("notes"):
        spec_pairs.append(("Broker Notes", specs.get("notes")))

    table_shape = slide.shapes.add_table(
        len(spec_pairs) + 1,
        2,
        Inches(0.6),
        Inches(1.4),
        Inches(5.6),
        Inches(4.6),
    )
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.4)

    for col_idx, h_text in enumerate(["Specification", "Details"]):
        c = table.cell(0, col_idx)
        c.fill.solid()
        c.fill.fore_color.rgb = DARK_NAVY
        p = c.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WARM_WHITE
        p.font.name = "Century Gothic"

    for idx, (label, val) in enumerate(spec_pairs, start=1):
        c0 = table.cell(idx, 0)
        c1 = table.cell(idx, 1)
        row_bg = GOLD_TINT if label == "Broker Notes" else (LIGHT_BG if idx % 2 == 0 else WARM_WHITE)
        c0.fill.solid()
        c0.fill.fore_color.rgb = row_bg
        c1.fill.solid()
        c1.fill.fore_color.rgb = row_bg

        p0 = c0.text_frame.paragraphs[0]
        p0.text = label
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.font.name = "Century Gothic"

        p1 = c1.text_frame.paragraphs[0]
        p1.text = str(val)
        p1.font.size = Pt(10)
        p1.font.bold = label in ("Gross Leasable Area", "Asking Lease Rate")
        p1.font.color.rgb = SLATE_TEXT
        p1.font.name = "Century Gothic"

    image_added = False
    thumb_path = slide_item.get("localThumbnailPath")
    if thumb_path and os.path.exists(thumb_path):
        try:
            slide.shapes.add_picture(
                thumb_path, Inches(6.5), Inches(1.4), width=Inches(6.2)
            )
            image_added = True
        except Exception:
            image_added = False

    if not image_added:
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(1.4), Inches(6.2), Inches(4.6)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = GOLD

        tb = slide.shapes.add_textbox(Inches(6.7), Inches(1.6), Inches(5.8), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = f"SOURCE DECK: {slide_item.get('deckFileName', '')} (Slide #{slide_item.get('slideNumber', 1)})"
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.font.name = "Century Gothic"

        raw_excerpt = (slide_item.get("rawText") or "").strip()[:550]
        p1 = tf.add_paragraph()
        p1.text = raw_excerpt or "Property visual and technical layout from master deck."
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = SLATE_TEXT
        p1.font.name = "Century Gothic"
        p1.space_before = Pt(10)


def add_credentials_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_solid_bg(slide, prs, WARM_WHITE)
    add_header_banner(
        slide,
        prs,
        "Advisory Scope & Capabilities",
        "PRIME Philippines — Industrial & Logistics Practice",
    )

    pillars = [
        (
            "1. Strategic Site Selection",
            "Comprehensive supply chain & corridor mapping across CALABARZON, Central Luzon, and Metro Manila.",
        ),
        (
            "2. Technical Due Diligence",
            "Verification of floor load capacity, clear apex height, power redundancy, PEZA accreditation, and truck turning radii.",
        ),
        (
            "3. Commercial Structuring",
            "Benchmark lease rate negotiation, escalation caps, rent-free fit-out periods, and build-to-suit (BTS) structuring.",
        ),
    ]

    for idx, (heading, body) in enumerate(pillars):
        left = Inches(0.6 + idx * 4.1)
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, Inches(1.7), Inches(3.85), Inches(4.5)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_BG
        box.line.color.rgb = GOLD

        tb = slide.shapes.add_textbox(left + Inches(0.25), Inches(1.95), Inches(3.35), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = heading
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = NAVY
        p0.font.name = "Century Gothic"

        p1 = tf.add_paragraph()
        p1.text = body
        p1.font.size = Pt(11)
        p1.font.color.rgb = SLATE_TEXT
        p1.font.name = "Century Gothic"
        p1.space_before = Pt(14)


def add_closing_slide(prs, proposal):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_solid_bg(slide, prs, NAVY)

    tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "THANK YOU"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = GOLD
    p0.font.name = "Century Gothic"

    p1 = tf.add_paragraph()
    broker_name = proposal.get("brokerName") or "PRIME Industrial Advisory Team"
    broker_email = proposal.get("brokerEmail") or "industrial@primephilippines.com"
    broker_phone = proposal.get("brokerPhone") or "+63 2 8888 0000"
    p1.text = f"{broker_name}\n{broker_email}   |   {broker_phone}\nwww.primephilippines.com"
    p1.font.size = Pt(14)
    p1.font.color.rgb = WARM_WHITE
    p1.font.name = "Century Gothic"
    p1.space_before = Pt(20)

    p2 = tf.add_paragraph()
    p2.text = "CONFIDENTIALITY NOTICE: This proposal is prepared exclusively for the intended recipient. Availability and commercial terms are subject to final lessor confirmation."
    p2.font.size = Pt(9)
    p2.font.color.rgb = MUTED_TEXT
    p2.font.name = "Century Gothic"
    p2.space_before = Pt(36)


def compile_manifest(manifest_path, output_path):
    with open(manifest_path, "r", encoding="utf-8") as f:
        manifest = json.load(f)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    proposal = manifest.get("proposal") or {}
    slides = manifest.get("slides") or []
    comparison_rows = manifest.get("comparisonRows") or []

    add_cover_slide(prs, proposal)

    if manifest.get("includeSummaryTable") and comparison_rows:
        add_summary_table_slide(prs, comparison_rows, proposal)

    for idx, slide_item in enumerate(slides, start=1):
        add_property_slide(prs, slide_item, idx)

    if manifest.get("includeCredentials"):
        add_credentials_slide(prs)

    add_closing_slide(prs, proposal)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)

    return {
        "outputPath": output_path,
        "slideCount": len(prs.slides),
        "engine": "python-pptx-prime-master",
    }


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(json.dumps({"error": "Usage: pptx_engine.py compile <manifest.json> <output.pptx>"}))
        sys.exit(1)

    command = sys.argv[1]
    if command == "compile":
        result = compile_manifest(sys.argv[2], sys.argv[3])
        print(json.dumps(result))
    else:
        print(json.dumps({"error": f"Unknown command: {command}"}))
        sys.exit(1)
